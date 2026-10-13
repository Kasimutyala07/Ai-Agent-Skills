
from pathlib import Path
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt


def create_presentation(title, slides):
    """
    Create a PowerPoint presentation.

    title: Presentation title
    slides: List of dictionaries containing
            slide titles and bullet points
    """
    output_dir = Path(__file__).resolve().parent / "output"
    output_dir.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Created by AI Agent"

    # Content slides
    for item in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = item["title"]

        body = slide.placeholders[1].text_frame
        body.clear()

        for index, point in enumerate(item["points"]):
            paragraph = (
                body.paragraphs[0]
                if index == 0
                else body.add_paragraph()
            )
            paragraph.text = str(point)
            paragraph.level = 0
            paragraph.font.size = Pt(24)

    # Use a unique filename for every presentation
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_title = "".join(
        c for c in title if c.isalnum() or c in " _-"
    ).strip()[:50] or "Presentation"

    file_path = output_dir / f"{safe_title}_{timestamp}.pptx"
    prs.save(file_path)

    return str(file_path)
