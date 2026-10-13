
from pathlib import Path
from docx import Document
from docx.shared import Pt


def create_word_document(title, sections):
    """
    Create a Word document from a title
    and a dictionary of section headings and content.
    """
    output_dir = Path(__file__).resolve().parent / "output"
    output_dir.mkdir(exist_ok=True)

    doc = Document()

    doc.add_heading(title, level=0)

    for heading, content in sections.items():
        doc.add_heading(heading, level=1)
        doc.add_paragraph(content)

    # Set a readable default font size
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    file_path = output_dir / "AI_Report.docx"
    doc.save(file_path)

    return str(file_path)
