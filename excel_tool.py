
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_excel(title, sheets):
    output_dir = Path(__file__).resolve().parent / "output"
    output_dir.mkdir(exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_data in sheets:
        ws = wb.create_sheet(
            title=sheet_data["name"][:31]
        )

        headers = sheet_data["headers"]
        rows = sheet_data.get("rows", [])

        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="217346"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row in rows:
            ws.append(row)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            longest = max(
                len(str(cell.value or ""))
                for cell in column
            )
            ws.column_dimensions[letter].width = min(
                longest + 3, 40
            )

    safe_title = "".join(
        c for c in title if c.isalnum() or c in " _-"
    ).strip()[:40] or "Workbook"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"{safe_title}_{timestamp}.xlsx"

    wb.save(path)
    return str(path)
